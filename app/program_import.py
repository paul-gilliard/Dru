"""Import de programmes coach depuis Excel / CSV / PDF / images via Claude.

Jobs stockés sous instance/program_imports/<id>/ (fichier(s) + meta.json).
Les captures d'écran sont acceptées, y compris plusieurs images pour un
même programme (elles sont envoyées dans l'ordre à Claude).
"""
from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
import uuid
from datetime import datetime, timezone
from typing import Any

from app.models import MUSCLE_GROUPS

DOC_EXTENSIONS = {'.xlsx', '.xls', '.csv', '.pdf'}
IMAGE_MEDIA_TYPES = {
    '.png': 'image/png',
    '.jpg': 'image/jpeg',
    '.jpeg': 'image/jpeg',
    '.jfif': 'image/jpeg',
    '.webp': 'image/webp',
    '.gif': 'image/gif',
    '.bmp': 'image/bmp',
    '.tif': 'image/tiff',
    '.tiff': 'image/tiff',
    '.heic': 'image/heic',
    '.heif': 'image/heif',
}
IMAGE_EXTENSIONS = set(IMAGE_MEDIA_TYPES)
ALLOWED_EXTENSIONS = DOC_EXTENSIONS | IMAGE_EXTENSIONS
# Formats acceptés nativement par l'API Anthropic (les autres sont convertis).
CLAUDE_IMAGE_MEDIA = {'image/jpeg', 'image/png', 'image/gif', 'image/webp'}

MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_TOTAL_BYTES = 40 * 1024 * 1024
MAX_IMAGES = 12
# Au-delà, Claude n'y gagne rien et le coût de tokens explose.
MAX_IMAGE_EDGE = 1800
# Sonnet 4 (20250514) retiré côté Anthropic — défaut = Sonnet 5.5 (vision + docs).
CLAUDE_MODEL = os.environ.get('ANTHROPIC_PROGRAM_IMPORT_MODEL', 'claude-sonnet-5-5')

MUSCLE_ALIASES = {
    'abdos': 'ABDOS', 'abs': 'ABDOS', 'abdominaux': 'ABDOS', 'core': 'ABDOS',
    'adducteur': 'ADDUCTEUR', 'adducteurs': 'ADDUCTEUR',
    'avant-bras': 'AVANT-BRAS', 'avant bras': 'AVANT-BRAS', 'forearm': 'AVANT-BRAS',
    'biceps': 'BICEPS',
    'dos': 'DOS', 'back': 'DOS', 'lats': 'DOS', 'dorsal': 'DOS', 'dorsaux': 'DOS',
    'epaules': 'EPAULES', 'épaule': 'EPAULES', 'épaules': 'EPAULES',
    'deltoide': 'EPAULES', 'deltoïdes': 'EPAULES',
    'deltoides': 'EPAULES', 'shoulder': 'EPAULES', 'shoulders': 'EPAULES',
    'fessiers': 'FESSIERS', 'fessier': 'FESSIERS', 'glutes': 'FESSIERS',
    'ischio': 'ISCHIO', 'ischios': 'ISCHIO', 'hamstring': 'ISCHIO', 'hamstrings': 'ISCHIO',
    'legs': 'LEGS', 'jambes': 'LEGS', 'leg': 'LEGS',
    'mollet': 'MOLLET', 'mollets': 'MOLLET', 'calves': 'MOLLET', 'calf': 'MOLLET',
    'pec': 'PEC', 'pectoraux': 'PEC', 'pectoral': 'PEC', 'chest': 'PEC', 'pecto': 'PEC',
    'quad': 'QUAD', 'quads': 'QUAD', 'quadriceps': 'QUAD',
    'triceps': 'TRICEPS',
}

DAY_ALIASES = {
    'lundi': 0, 'monday': 0, 'lun': 0, 'mon': 0,
    'mardi': 1, 'tuesday': 1, 'mar': 1, 'tue': 1,
    'mercredi': 2, 'wednesday': 2, 'mer': 2, 'wed': 2,
    'jeudi': 3, 'thursday': 3, 'jeu': 3, 'thu': 3,
    'vendredi': 4, 'friday': 4, 'ven': 4, 'fri': 4,
    'samedi': 5, 'saturday': 5, 'sam': 5, 'sat': 5,
    'dimanche': 6, 'sunday': 6, 'dim': 6, 'sun': 6,
}

SYSTEM_PROMPT = """Tu es un expert musculation qui convertit des programmes d'entraînement (Excel, CSV, PDF,
photos ou captures d'écran) vers le format Farmness. Extrais UNIQUEMENT les séances de musculation
(ignore diètes, journaux, métriques).

Règles:
- Si plusieurs images sont fournies, elles forment UN SEUL programme découpé en morceaux : lis-les dans l'ordre,
  fusionne les séances, et ne duplique pas une séance qui apparaît à cheval sur deux images.
- Sur une image, respecte la structure visuelle (colonnes, blocs de couleur, en-têtes de séance).
  Si un texte est illisible, ne l'invente pas : mets-le dans warnings.
- day_of_week: 0=lundi … 6=dimanche. Si "Jour 1/2/3" sans jour nommé, mappe 0,1,2… en gardant l'ordre.
- Muscle Farmness UNIQUEMENT parmi: ABDOS, ADDUCTEUR, AVANT-BRAS, BICEPS, DOS, EPAULES, FESSIERS, ISCHIO, LEGS, MOLLET, PEC, QUAD, TRICEPS.
  Mappe Pectoraux→PEC, Deltoïde/Épaules→EPAULES, etc. Si absent, déduis du mouvement.
- Conserve reps/rest tels quels (ex: "6-8 + 15 + 15", "4'+2'+2'", "MAX-MAX-MAX").
- sets = nombre de séries principales (entier) ; si ambigu, estime raisonnablement.
- remark = notes d'exécution / amplitude du coach.
- Remplis le muscle group "fill-down" si la cellule est vide mais le bloc précédent l'indique.
- Ignore échauffements génériques sans mouvement précis (ex: "vélo 5 min") sauf s'ils sont clairement un exercice nommé.
- program_name: titre du fichier/onglet ou "Programme importé".
- warnings: ambiguïtés pour le coach.

Réponds UNIQUEMENT avec un JSON valide (pas de markdown) de forme:
{
  "program_name": "string",
  "sessions": [
    {
      "day_of_week": 0,
      "session_name": "string",
      "exercises": [
        {
          "name": "string",
          "sets": 3,
          "reps": "string|null",
          "rest": "string|null",
          "muscle": "PEC",
          "remark": "string|null",
          "intensification": "string|null"
        }
      ]
    }
  ],
  "warnings": ["string"]
}
"""


def _imports_root() -> str:
    basedir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    root = os.path.join(basedir, 'instance', 'program_imports')
    os.makedirs(root, exist_ok=True)
    return root


def _assert_safe_job_id(job_id: str) -> str:
    jid = (job_id or '').strip()
    if not re.fullmatch(r'[a-f0-9]{32}', jid):
        raise ValueError('import id invalide')
    return jid


def _job_dir(job_id: str) -> str:
    jid = _assert_safe_job_id(job_id)
    path = os.path.join(_imports_root(), jid)
    os.makedirs(path, exist_ok=True)
    return path


def _meta_path(job_id: str) -> str:
    return os.path.join(_job_dir(job_id), 'meta.json')


def _file_path(job_id: str, filename: str) -> str:
    safe = re.sub(r'[^\w.\-]+', '_', filename)[:120] or 'upload.bin'
    return os.path.join(_job_dir(job_id), safe)


def load_job(job_id: str) -> dict | None:
    try:
        jid = _assert_safe_job_id(job_id)
    except ValueError:
        return None
    path = os.path.join(_imports_root(), jid, 'meta.json')
    if not os.path.isfile(path):
        return None
    with open(path, 'r', encoding='utf-8') as fh:
        return json.load(fh)


def save_job(meta: dict) -> None:
    meta['updated_at'] = datetime.now(timezone.utc).isoformat()
    with open(_meta_path(meta['id']), 'w', encoding='utf-8') as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=2)


def _fold(s: str) -> str:
    table = str.maketrans({
        'à': 'a', 'â': 'a', 'ä': 'a', 'á': 'a', 'ã': 'a',
        'é': 'e', 'è': 'e', 'ê': 'e', 'ë': 'e',
        'î': 'i', 'ï': 'i', 'í': 'i', 'ì': 'i',
        'ô': 'o', 'ö': 'o', 'ó': 'o', 'ò': 'o',
        'û': 'u', 'ü': 'u', 'ú': 'u', 'ù': 'u',
        'ÿ': 'y', 'ç': 'c', 'ñ': 'n',
    })
    return (s or '').lower().translate(table)


def normalize_muscle(raw: str | None) -> str | None:
    if not raw:
        return None
    s = (raw or '').strip()
    if s.upper() in MUSCLE_GROUPS:
        return s.upper()
    key = _fold(s).replace('_', ' ').strip()
    if key in MUSCLE_ALIASES:
        return MUSCLE_ALIASES[key]
    for alias, code in MUSCLE_ALIASES.items():
        if alias in key or key in alias:
            return code
    return None


def _ext_of(filename: str) -> str:
    return os.path.splitext(filename or '')[1].lower()


def detect_workbook_sheets(file_bytes: bytes) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def detect_pdf_page_count(file_bytes: bytes) -> int:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    return len(reader.pages)


def sheet_to_tsv(file_bytes: bytes, sheet_name: str | None = None, max_rows: int = 400) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        name = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
        ws = wb[name]
        lines = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                lines.append(f'… (tronqué après {max_rows} lignes)')
                break
            cells = []
            for cell in row:
                if cell is None:
                    cells.append('')
                else:
                    cells.append(str(cell).replace('\t', ' ').replace('\n', ' ').strip())
            # Skip fully empty rows
            if any(cells):
                lines.append('\t'.join(cells))
        return f'[Onglet: {name}]\n' + '\n'.join(lines)
    finally:
        wb.close()


def csv_to_text(file_bytes: bytes, max_rows: int = 500) -> str:
    text = None
    for enc in ('utf-8-sig', 'utf-8', 'latin-1', 'cp1252'):
        try:
            text = file_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = file_bytes.decode('utf-8', errors='replace')
    reader = csv.reader(io.StringIO(text))
    lines = []
    for i, row in enumerate(reader):
        if i >= max_rows:
            lines.append(f'… (tronqué après {max_rows} lignes)')
            break
        lines.append('\t'.join((c or '').replace('\t', ' ').replace('\n', ' ').strip() for c in row))
    return '\n'.join(lines)


def pdf_page_text(file_bytes: bytes, page_index: int | None = None, max_chars: int = 60000) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    if page_index is not None:
        if page_index < 0 or page_index >= len(reader.pages):
            raise ValueError('page invalide')
        pages = [reader.pages[page_index]]
        indices = [page_index]
    else:
        pages = list(reader.pages)
        indices = list(range(len(pages)))
    for idx, page in zip(indices, pages):
        raw = (page.extract_text() or '').strip()
        parts.append(f'[Page {idx + 1}]\n{raw}')
    out = '\n\n'.join(parts)
    return out[:max_chars]


def _sniff_image_media(raw: bytes) -> str | None:
    """Détecte le type réel : les screens arrivent souvent avec une extension fausse."""
    if raw[:8] == b'\x89PNG\r\n\x1a\n':
        return 'image/png'
    if raw[:3] == b'\xff\xd8\xff':
        return 'image/jpeg'
    if raw[:6] in (b'GIF87a', b'GIF89a'):
        return 'image/gif'
    if raw[:4] == b'RIFF' and raw[8:12] == b'WEBP':
        return 'image/webp'
    if raw[:2] == b'BM':
        return 'image/bmp'
    if raw[:4] in (b'II*\x00', b'MM\x00*'):
        return 'image/tiff'
    if raw[4:12] in (b'ftypheic', b'ftypheix', b'ftyphevc', b'ftypmif1', b'ftypmsf1'):
        return 'image/heic'
    return None


def normalize_image(raw: bytes, ext: str) -> tuple[bytes, str]:
    """Ramène n'importe quelle image à un format lisible par Claude, redimensionnée."""
    # L'extension seule ne prouve rien : on ne se rabat que sur la signature binaire.
    media = _sniff_image_media(raw)
    try:
        from PIL import Image
        try:
            import pillow_heif
            pillow_heif.register_heif_opener()
        except Exception:
            pass
        img = Image.open(io.BytesIO(raw))
        img.load()
        if img.mode not in ('RGB', 'L'):
            img = img.convert('RGB')
        width, height = img.size
        longest = max(width, height)
        if longest > MAX_IMAGE_EDGE:
            scale = MAX_IMAGE_EDGE / float(longest)
            img = img.resize((max(1, int(width * scale)), max(1, int(height * scale))), Image.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, format='JPEG', quality=88, optimize=True)
        return buf.getvalue(), 'image/jpeg'
    except Exception:
        if media in CLAUDE_IMAGE_MEDIA:
            return raw, media
        raise ValueError(
            'Image illisible par le serveur — convertis-la en JPEG ou PNG et réessaie'
        )


def create_upload_job(
    *,
    coach_id: int,
    files: list[tuple[str, bytes]] | None = None,
    filename: str | None = None,
    file_bytes: bytes | None = None,
) -> dict:
    """Crée un job d'import. `files` = [(nom, contenu)] ; multi-fichiers réservé aux images."""
    items: list[tuple[str, bytes]] = list(files or [])
    if not items and filename is not None and file_bytes is not None:
        items = [(filename, file_bytes)]
    items = [(name, blob) for name, blob in items if blob]
    if not items:
        raise ValueError('Aucun fichier reçu')

    total = 0
    for name, blob in items:
        if len(blob) > MAX_UPLOAD_BYTES:
            raise ValueError(f'"{name}" trop volumineux (max {MAX_UPLOAD_BYTES // (1024 * 1024)} Mo par fichier)')
        total += len(blob)
    if total > MAX_TOTAL_BYTES:
        raise ValueError(f'Import trop volumineux (max {MAX_TOTAL_BYTES // (1024 * 1024)} Mo au total)')

    exts: list[str] = []
    for name, blob in items:
        ext = _ext_of(name)
        if ext not in ALLOWED_EXTENSIONS:
            # Les captures partagées arrivent parfois sans extension fiable.
            sniffed = _sniff_image_media(blob)
            ext = next((e for e, m in IMAGE_MEDIA_TYPES.items() if m == sniffed), '') if sniffed else ''
        if ext not in ALLOWED_EXTENSIONS:
            raise ValueError(
                f'Format non géré pour "{name}". Accepté : Excel (.xlsx/.xls), CSV, PDF, images (JPG, PNG, HEIC…)'
            )
        exts.append(ext)

    images_only = all(ext in IMAGE_EXTENSIONS for ext in exts)
    if len(items) > 1 and not images_only:
        raise ValueError('Plusieurs fichiers à la fois : uniquement pour des images / captures d\'écran')
    if images_only and len(items) > MAX_IMAGES:
        raise ValueError(f'Maximum {MAX_IMAGES} images par import')

    job_id = uuid.uuid4().hex
    stored_files: list[dict] = []
    sheets: list[str] = []
    page_count: int | None = None

    if images_only:
        for idx, (name, blob) in enumerate(items):
            data, media_type = normalize_image(blob, exts[idx])
            suffix = '.jpg' if media_type == 'image/jpeg' else (
                '.png' if media_type == 'image/png' else exts[idx] or '.img'
            )
            stored = f'page_{idx + 1}{suffix}'
            with open(_file_path(job_id, stored), 'wb') as fh:
                fh.write(data)
            stored_files.append({
                'name': name,
                'stored': stored,
                'media_type': media_type,
                'bytes': len(data),
            })
        kind = 'images'
        ext = exts[0]
        display_name = items[0][0] if len(items) == 1 else f'{len(items)} images'
    else:
        name, blob = items[0]
        ext = exts[0]
        stored = f'upload{ext}'
        with open(_file_path(job_id, stored), 'wb') as fh:
            fh.write(blob)
        stored_files.append({'name': name, 'stored': stored, 'media_type': None, 'bytes': len(blob)})
        display_name = name
        if ext in ('.xlsx', '.xls'):
            kind = 'workbook'
            try:
                sheets = detect_workbook_sheets(blob)
            except Exception as exc:
                raise ValueError(f'Excel illisible : {exc}') from exc
        elif ext == '.csv':
            kind = 'csv'
        else:
            kind = 'pdf'
            try:
                page_count = detect_pdf_page_count(blob)
            except Exception as exc:
                raise ValueError(f'PDF illisible : {exc}') from exc

    meta = {
        'id': job_id,
        'coach_id': int(coach_id),
        'filename': display_name,
        'stored_name': stored_files[0]['stored'],
        'ext': ext,
        'kind': kind,
        'files': stored_files,
        'image_count': len(stored_files) if kind == 'images' else 0,
        'sheets': sheets,
        'page_count': page_count,
        'status': 'uploaded',
        'selected_sheet': None,
        'selected_page': None,
        'hint': None,
        'draft': None,
        'match_items': [],
        'warnings': [],
        'program_id': None,
        'created_at': datetime.now(timezone.utc).isoformat(),
        'updated_at': datetime.now(timezone.utc).isoformat(),
    }
    save_job(meta)
    return meta


def job_kind(meta: dict) -> str:
    kind = (meta.get('kind') or '').strip()
    if kind:
        return kind
    ext = meta.get('ext') or ''
    if ext in IMAGE_EXTENSIONS:
        return 'images'
    if ext in ('.xlsx', '.xls'):
        return 'workbook'
    if ext == '.csv':
        return 'csv'
    return 'pdf'


def read_job_bytes(meta: dict) -> bytes:
    path = _file_path(meta['id'], meta['stored_name'])
    with open(path, 'rb') as fh:
        return fh.read()


def read_job_images(meta: dict) -> list[dict]:
    out: list[dict] = []
    for entry in meta.get('files') or []:
        path = _file_path(meta['id'], entry.get('stored') or '')
        if not os.path.isfile(path):
            continue
        with open(path, 'rb') as fh:
            raw = fh.read()
        out.append({
            'name': entry.get('name'),
            'media_type': entry.get('media_type') or _sniff_image_media(raw) or 'image/jpeg',
            'bytes': raw,
        })
    return out


def build_extract_payload(meta: dict, *, sheet: str | None, page: int | None) -> dict[str, Any]:
    """Prépare le contenu à envoyer à Claude (texte, PDF natif ou images)."""
    kind = job_kind(meta)
    hint_parts = []

    if kind == 'images':
        images = read_job_images(meta)
        if not images:
            raise ValueError('images introuvables pour cet import')
        return {
            'mode': 'images',
            'selected_sheet': None,
            'selected_page': None,
            'text': None,
            'file_bytes': None,
            'media_type': None,
            'images': images,
        }

    file_bytes = read_job_bytes(meta)
    ext = meta['ext']

    if ext in ('.xlsx', '.xls'):
        sheets = meta.get('sheets') or []
        chosen = sheet if sheet in sheets else (sheets[0] if sheets else None)
        text = sheet_to_tsv(file_bytes, chosen)
        return {
            'mode': 'text',
            'selected_sheet': chosen,
            'selected_page': None,
            'text': text,
            'file_bytes': None,
            'media_type': None,
        }

    if ext == '.csv':
        return {
            'mode': 'text',
            'selected_sheet': None,
            'selected_page': None,
            'text': csv_to_text(file_bytes),
            'file_bytes': None,
            'media_type': None,
        }

    # PDF : préférer document natif Anthropic ; si page isolée, tenter texte d'abord
    if page is not None:
        try:
            text = pdf_page_text(file_bytes, page)
            if len(text.strip()) > 80:
                return {
                    'mode': 'text',
                    'selected_sheet': None,
                    'selected_page': page,
                    'text': text,
                    'file_bytes': None,
                    'media_type': None,
                }
        except Exception:
            hint_parts.append('extraction texte page échouée → PDF natif')

    return {
        'mode': 'pdf',
        'selected_sheet': None,
        'selected_page': page,
        'text': None,
        'file_bytes': file_bytes,
        'media_type': 'application/pdf',
        'note': '; '.join(hint_parts) if hint_parts else None,
    }


def _extract_json_object(raw: str) -> dict:
    raw = (raw or '').strip()
    if raw.startswith('```'):
        raw = re.sub(r'^```(?:json)?\s*', '', raw)
        raw = re.sub(r'\s*```$', '', raw)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start = raw.find('{')
        end = raw.rfind('}')
        if start >= 0 and end > start:
            return json.loads(raw[start:end + 1])
        raise


def normalize_draft(raw: dict) -> dict:
    name = (raw.get('program_name') or 'Programme importé').strip()[:128] or 'Programme importé'
    sessions_out = []
    used_days: set[int] = set()
    for i, sess in enumerate(raw.get('sessions') or []):
        if not isinstance(sess, dict):
            continue
        dow = sess.get('day_of_week')
        try:
            dow = int(dow)
        except (TypeError, ValueError):
            # Essaye nom de session
            label = _fold(str(sess.get('session_name') or ''))
            dow = None
            for key, val in DAY_ALIASES.items():
                if key in label:
                    dow = val
                    break
            if dow is None:
                dow = i % 7
        dow = max(0, min(6, int(dow)))
        # Collision : décale au prochain jour libre
        if dow in used_days:
            for candidate in range(7):
                if candidate not in used_days:
                    dow = candidate
                    break
        used_days.add(dow)
        session_name = (sess.get('session_name') or f'Séance jour {dow + 1}').strip()[:128]
        exercises = []
        last_muscle = None
        for ex in sess.get('exercises') or []:
            if not isinstance(ex, dict):
                continue
            ename = (ex.get('name') or '').strip()[:192]
            if not ename:
                continue
            muscle = normalize_muscle(ex.get('muscle')) or last_muscle
            if muscle:
                last_muscle = muscle
            sets = ex.get('sets')
            try:
                sets = int(sets) if sets is not None and str(sets).strip() != '' else None
            except (TypeError, ValueError):
                sets = None
            exercises.append({
                'name': ename,
                'sets': sets,
                'reps': (str(ex['reps'])[:64] if ex.get('reps') is not None else None),
                'rest': (str(ex['rest'])[:64] if ex.get('rest') is not None else None),
                'muscle': muscle,
                'remark': (str(ex['remark'])[:2000] if ex.get('remark') else None),
                'intensification': (str(ex['intensification'])[:64] if ex.get('intensification') else None),
            })
        if exercises:
            sessions_out.append({
                'day_of_week': dow,
                'session_name': session_name,
                'exercises': exercises,
            })
    warnings = [str(w) for w in (raw.get('warnings') or []) if w][:20]
    return {'program_name': name, 'sessions': sessions_out, 'warnings': warnings}


def call_claude_parse(*, extract: dict, hint: str | None, filename: str) -> dict:
    """Appelle Anthropic. Overridable en tests via monkeypatch."""
    api_key = (os.environ.get('ANTHROPIC_API_KEY') or '').strip()
    if not api_key:
        raise RuntimeError("L'import intelligent est temporairement indisponible.")

    import anthropic

    user_bits: list[Any] = []
    coach_hint = (hint or '').strip()
    preamble = (
        f"Fichier source : {filename}\n"
        f"Consignes coach (optionnel) : {coach_hint or '(aucune)'}\n"
        "Extrais le programme musculation au format JSON demandé."
    )

    if extract['mode'] == 'images' and extract.get('images'):
        images = extract['images']
        user_bits.append({
            'type': 'text',
            'text': (
                f"{preamble}\n\n"
                f"{len(images)} image(s) fournie(s), dans l'ordre. "
                "Elles décrivent un seul et même programme."
            ),
        })
        for idx, img in enumerate(images):
            user_bits.append({'type': 'text', 'text': f'Image {idx + 1}/{len(images)} :'})
            user_bits.append({
                'type': 'image',
                'source': {
                    'type': 'base64',
                    'media_type': img['media_type'],
                    'data': base64.standard_b64encode(img['bytes']).decode('ascii'),
                },
            })
    elif extract['mode'] == 'pdf' and extract.get('file_bytes'):
        b64 = base64.standard_b64encode(extract['file_bytes']).decode('ascii')
        user_bits.append({
            'type': 'document',
            'source': {
                'type': 'base64',
                'media_type': 'application/pdf',
                'data': b64,
            },
        })
        user_bits.append({'type': 'text', 'text': preamble})
    else:
        text = extract.get('text') or ''
        user_bits.append({
            'type': 'text',
            'text': f"{preamble}\n\n--- CONTENU ---\n{text}",
        })

    model = (os.environ.get('ANTHROPIC_PROGRAM_IMPORT_MODEL') or CLAUDE_MODEL).strip()
    client = anthropic.Anthropic(api_key=api_key)
    create_kwargs = {
        'model': model,
        'max_tokens': 8192,
        'system': SYSTEM_PROMPT,
        'messages': [{'role': 'user', 'content': user_bits}],
    }
    # Sonnet 5.x : effort bas = peu de thinking, JSON plus fiable / moins cher
    if model.startswith('claude-sonnet-5') or model.startswith('claude-opus-5'):
        create_kwargs['output_config'] = {'effort': 'low'}
    try:
        msg = client.messages.create(**create_kwargs)
    except TypeError:
        create_kwargs.pop('output_config', None)
        msg = client.messages.create(**create_kwargs)
    parts = []
    for block in msg.content:
        if getattr(block, 'type', None) == 'text':
            parts.append(block.text)
    raw = _extract_json_object('\n'.join(parts))
    return normalize_draft(raw)


def build_match_items(draft: dict) -> list[dict]:
    """Une entrée unique par nom d'exercice source (réutilisé dans plusieurs séances)."""
    seen: dict[str, dict] = {}
    order: list[str] = []
    for si, sess in enumerate(draft.get('sessions') or []):
        for ei, ex in enumerate(sess.get('exercises') or []):
            name = (ex.get('name') or '').strip()
            if not name:
                continue
            key = _fold(name)
            if key not in seen:
                seen[key] = {
                    'key': f'ex_{len(order)}',
                    'source_name': name,
                    'muscle': ex.get('muscle'),
                    'occurrences': [],
                    'candidates': [],
                    'suggested': None,
                    'resolution': None,
                }
                order.append(key)
            seen[key]['occurrences'].append({'session_idx': si, 'exercise_idx': ei})
            if not seen[key]['muscle'] and ex.get('muscle'):
                seen[key]['muscle'] = ex.get('muscle')
    return [seen[k] for k in order]


def public_job_view(meta: dict, *, include_draft: bool = True) -> dict:
    out = {
        'id': meta['id'],
        'filename': meta.get('filename'),
        'ext': meta.get('ext'),
        'kind': job_kind(meta),
        'image_count': int(meta.get('image_count') or 0),
        'files': [{'name': f.get('name')} for f in (meta.get('files') or [])],
        'sheets': meta.get('sheets') or [],
        'page_count': meta.get('page_count'),
        'status': meta.get('status'),
        'selected_sheet': meta.get('selected_sheet'),
        'selected_page': meta.get('selected_page'),
        'hint': meta.get('hint'),
        'warnings': meta.get('warnings') or [],
        'program_id': meta.get('program_id'),
        'match_items': meta.get('match_items') or [],
    }
    if include_draft:
        out['draft'] = meta.get('draft')
    return out
